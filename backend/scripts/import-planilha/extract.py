#!/usr/bin/env python3
"""Importação única da planilha "Controle Financeiro" (abas AGO26-ATUAL e JUN-24-JUL26).

Etapa 1 (este script): lê o .xlsx (valores + comentários das células) e gera
  - dataset.json : dados estruturados, prontos para o import.ts
  - review.md    : relatório de revisão (contagens, divergências, itens sinalizados)
Nada é gravado no banco aqui.

Uso: python extract.py "C:/caminho/Controle Financeiro.xlsx" [pasta_de_saida]
"""
import collections
import io
import json
import os
import re
import sys
import unicodedata
from decimal import Decimal

import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'out')
os.makedirs(OUT, exist_ok=True)

TODAY_MONTH = '2026-10'     # mês corrente: a partir dele nada está pago
FIRST_MONTH = '2026-01'     # início do histórico a importar

# ---------------------------------------------------------------- configuração
CARDS = {
    'Itaú Visa':        dict(bestPurchaseDay=8, dueDay=15),
    'ML Visa':          dict(bestPurchaseDay=9, dueDay=14),
    'Bradesco Amazon':  dict(bestPurchaseDay=2, dueDay=15),
    'C6 Mastercard':    dict(bestPurchaseDay=9, dueDay=15),
}
# aba -> colunas: tipo card | pix
SHEETS = [
    dict(name='JUN-24-JUL26', rows=range(21, 28),
         cols={'H': ('card', 'Itaú Visa'), 'I': ('card', 'ML Visa'), 'K': ('card', 'C6 Mastercard'),
               'L': ('card', 'Bradesco Amazon'), 'J': ('pix', None), 'O': ('pix', None)},
         fixed={'B': ('Aluguel + encargos', 'Moradia', 8), 'C': ('RCI Sandero (financiamento)', 'Carro - financiamento', 10),
                'D': ('Energia', 'Contas fixas', 10), 'E': ('Água', 'Contas fixas', 10),
                'F': ('Internet', 'Contas fixas', 15), 'G': ('Celular (plano)', 'Contas fixas', 15)},
         receber='Q', entradas='S'),
    dict(name='AGO26-ATUAL', rows=range(2, 28),
         cols={'G': ('card', 'Itaú Visa'), 'H': ('card', 'ML Visa'), 'I': ('card', 'C6 Mastercard'),
               'J': ('card', 'Bradesco Amazon'), 'M': ('pix', None)},
         fixed={'B': ('Aluguel + encargos', 'Moradia', 8), 'C': ('Energia', 'Contas fixas', 10),
                'D': ('Água', 'Contas fixas', 10), 'E': ('Internet', 'Contas fixas', 15),
                'F': ('Celular (plano)', 'Contas fixas', 15)},
         receber='O', entradas='Q'),
]
FUTURE_FIXED = {  # recorrências a partir do mês seguinte ao corrente
    'Aluguel + encargos': dict(amount=206482, category='Moradia', dueDay=8),
    'Internet': dict(amount=11990, category='Contas fixas', dueDay=15),
    'Celular (plano)': dict(amount=10400, category='Contas fixas', dueDay=15),
    'Energia': dict(amount=18000, category='Contas fixas', dueDay=10),
    'Água': dict(amount=9200, category='Contas fixas', dueDay=10),
}
SALARY_FUTURE_CENTS = 553070
SUBSCRIPTION_NAMES = ('prime', 'spotify', 'ml hbo', 'ml', 'netflix', 'wellhub', 'claude')

flags = []   # (tipo, local, texto)


def flag(kind, where, text):
    flags.append((kind, where, text))


# ---------------------------------------------------------------- utilidades
def ym_index(ym):
    y, m = ym.split('-')
    return int(y) * 12 + int(m) - 1


def ym_from_index(i):
    return f'{i // 12:04d}-{i % 12 + 1:02d}'


def add_months(ym, n):
    return ym_from_index(ym_index(ym) + n)


def to_cents_str(s):
    s = s.strip()
    neg = s.startswith('-')
    s = s.lstrip('-').replace('.', '').replace(',', '.')
    c = int((Decimal(s) * 100).to_integral_value())
    return -c if neg else c


def to_cents_float(v):
    return int(round(float(v) * 100))


def strip_accents(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def norm(s):
    return re.sub(r'\s+', ' ', strip_accents(s).lower()).strip()


def body(c):
    t = c.comment.text
    i = t.find('Comment:')
    return t[i + 8:].strip() if i >= 0 else t


NUM = re.compile(r'(?<![\w/,.])-?\d{1,3}(?:\.\d{3})*(?:,\d{1,2})?(?![\w/])|(?<![\w/,.])-?\d+(?:,\d{1,2})?(?![\w/])')
FRAC = re.compile(r'(?<!\d)(\d{1,2})/(\d{1,2})(?!\d)')


def pick_amount(nums, line):
    if len(nums) == 1:
        return nums[0]
    dec = [n for n in nums if ',' in n]
    if len(dec) == 1:
        return dec[0]
    pool = dec or nums
    return pool[0] if re.match(r'\s*-?\d', line) else pool[-1]


def clean_desc(s):
    s = re.sub(r'\(.*?\)', ' ', s)
    s = re.sub(NUM, ' ', s)
    s = re.sub(r'\b\dx\b', ' ', s)
    s = re.sub(r'(?i)\b(pago|ok|estornado)\b', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip(' -:+,.')
    return s


def parse_line(line):
    raw = line.strip().replace('\xa0', ' ')
    if not raw:
        return None
    ms = list(FRAC.finditer(raw))
    p = None
    rest = raw
    if ms:
        mm = ms[-1]  # "Parcelamento março/26 1/3 272,35": só o último n/t vale
        p = (int(mm[1]), int(mm[2]))
        rest = raw[:mm.start()] + ' ' + raw[mm.end():]
    nums = NUM.findall(re.sub(r'\(.*?\)', ' ', rest))
    low = raw.lower()
    return dict(raw=raw, p=p, nums=nums, rest=rest,
                credit=bool(re.match(r'(?i)\s*estorno', raw)),
                estornado='estornado' in low)


# ---------------------------------------------------------------- categorias
def category_for(desc):
    d = norm(desc)

    def has(*ws):
        return any(w in d for w in ws)

    if has('parcelamento marco'):
        return 'Parcelamento de fatura'
    if has('airbnb'):
        return 'Larissa'
    if has('emprestad'):
        return 'Empréstimo'
    if has('caixinha'):
        return 'Reserva'
    if has('cigarro', 'sabia', 'isqueiro'):
        return 'Cigarro'
    if has('ifood', '99 food', '99food'):
        return 'Delivery'
    if has('seguro', 'ipva', 'licenciamento'):
        return 'Carro - documentação'
    if has('filtro combustivel', 'bomba combustivel', 'autozone', 'oleo', 'coxim', 'kit manutencao', 'bico injetor',
           'bobina', 'axial', 'terminal', 'filtro ar', 'forro capo', 'shampoo carro', 'freio', 'junta escape',
           'bolsa carro', 'escapamento', 'arruela', 'carter', 'alinhamento', 'borracheiro', 'fusivel carro',
           'insulfime', 'lavagem carro', 'aditivo', 'swift', 'mao de obra', 'limpeza bico', 'corsa', 'chave impacto'):
        return 'Carro - manutenção'
    if has('gasolina', 'shell', 'etanol'):
        return 'Combustível'
    if d == 'ml' or has('prime', 'spotify', 'netflix', 'hbo', 'premiere', 'wellhub', 'claude', 'gamepass'):
        return 'Assinaturas'
    if has('farmacia', 'remedio', 'drogasil', 'dentista', 'psicolog', 'psiclog', 'dorflex'):
        return 'Saúde'
    if has('polipet', 'racao', 'veterinario', 'areia', 'manta gato'):
        return 'Pets'
    if has('barbeiro', 'barbearia', 'cabelo', 'nair'):
        return 'Cuidados pessoais'
    if has('presente') or d.startswith('pai ') or ' pai' in d or d.startswith('pix mae') or ' mae' in d \
            or d.startswith('mae'):
        return 'Presentes e família'
    if has('restaurante', 'outback', 'pizza', 'lanche', 'popeyes', 'bk', 'sorvete', 'acai', 'cinema', 'churrasco',
           'pipoca', 'almoco', 'domino', 'hotel', 'frango assado', 'kebab', 'galeria', 'psn', 'chiquinho', 'estrela'):
        return 'Lazer'
    if has('shibata', 'semar', 'maximo', 'colinas', 'quitanda', 'tenda', 'mercado', 'pao de a', 'cafe', 'porcaria',
           'oxxo', 'coca', 'gelo', 'agua coco', 'salgadinho', 'cacaushow', 'frazili', 'doce'):
        return 'Mercado'
    if has('padaria', 'pao', 'bolo', 'salgado'):
        return 'Padaria'
    if has('ferragem', 'sodimac', 'magalu', 'americanas', 'panela', 'encanador', 'tinta', 'parafuso', 'lampada',
           'saco lixo', 'rodo', 'desinfetante', 'sanduicheira', 'dolce gusto', 'climatizador', 'extensao',
           'esmerilhadeira', 'ferramenta', 'carregador', 'esquadro', 'maquina', 'havan', 'daiso', 'mop', 'fonte',
           'pipoqueira', 'repelente', 'emenda', 'peneira', 'cooler', 'quadro', 'suporte', 'camera', 'aliexpress',
           'netshoes', 'kit chave', 'macaneta', 'produto piscina', 'sapateira', 'bateria controle'):
        return 'Casa'
    return 'Outros'


CATEGORIES = [
    ('Moradia', 'EXPENSE'), ('Contas fixas', 'EXPENSE'), ('Mercado', 'EXPENSE'), ('Padaria', 'EXPENSE'),
    ('Delivery', 'EXPENSE'), ('Lazer', 'EXPENSE'), ('Combustível', 'EXPENSE'), ('Carro - manutenção', 'EXPENSE'),
    ('Carro - documentação', 'EXPENSE'), ('Carro - financiamento', 'EXPENSE'), ('Assinaturas', 'EXPENSE'),
    ('Saúde', 'EXPENSE'), ('Casa', 'EXPENSE'), ('Pets', 'EXPENSE'), ('Cuidados pessoais', 'EXPENSE'),
    ('Cigarro', 'EXPENSE'), ('Presentes e família', 'EXPENSE'), ('Larissa', 'EXPENSE'), ('Empréstimo', 'EXPENSE'),
    ('Reserva', 'EXPENSE'), ('Parcelamento de fatura', 'EXPENSE'), ('Outros', 'EXPENSE'),
    ('Salário', 'INCOME'), ('Larissa (entrada)', 'INCOME'), ('Saldo anterior', 'INCOME'), ('Extras', 'INCOME'),
]

# ---------------------------------------------------------------- leitura
wb = openpyxl.load_workbook(SRC)

card_items = []          # compras de cartão avulsas
pix_items = []           # débito/PIX avulsos
credits = []             # estornos/créditos de cartão
raw_plan_items = []      # itens n/total em cartão
plain_installments = []  # n/total pagos em PIX/cartão de terceiros
sheet_cells = {}         # (card, month) -> valor da célula
incomes = []
fixed_rows = []
receber = []
months_seen = []


def month_of(ws, r):
    return ws[f'A{r}'].value.strftime('%Y-%m')


for sh in SHEETS:
    ws = wb[sh['name']]
    for r in sh['rows']:
        ym = month_of(ws, r)
        months_seen.append(ym)
        for col, (kind, target) in sh['cols'].items():
            c = ws[f'{col}{r}']
            where = f"{sh['name']}!{col}{r} ({ym})"
            if kind == 'card' and c.value is not None:
                sheet_cells[(target, ym)] = to_cents_float(c.value)
            if not c.comment:
                continue
            label = None
            for ln in body(c).split('\n'):
                res = parse_line(ln)
                if not res:
                    continue
                raw = res['raw']
                low = norm(raw)
                if low.startswith('assinaturas total') or low == 'fixas':
                    continue
                if raw.startswith('-'):
                    flag('linha ignorada (começa com "-": provável item devolvido/descontado)', where, raw)
                    continue
                if kind == 'card' and 'entrada' in low and len(res['nums']) >= 3:
                    nums = [to_cents_str(n) for n in res['nums']]
                    fin = nums[0] - nums[1]
                    credits.append(dict(card=target, month=ym, amount=fin,
                                        description='Parcelamento da fatura (valor financiado)',
                                        anchor=None, where=where))
                    flag('parcelamento de fatura (crédito de %.2f)' % (fin / 100), where, raw)
                    continue
                if not res['nums']:
                    if raw.endswith(':'):
                        label = clean_desc(raw)
                    else:
                        flag('linha sem valor (ignorada)', where, raw)
                    continue
                if label and len(res['nums']) >= 2 and not clean_desc(res['rest']):
                    for n in res['nums']:
                        card_items.append(dict(card=target, month=ym, description=label,
                                               amount=to_cents_str(n), where=where))
                    label = None
                    continue
                label = None
                amt = to_cents_str(pick_amount(res['nums'], raw))
                if len(res['nums']) > 1 and not any(',' in n for n in res['nums']):
                    flag('mais de um número na linha (usei o primeiro/último)', where, raw)
                desc = clean_desc(res['rest']) or 'Sem descrição'
                if kind == 'pix' and res['p']:
                    plain_installments.append(dict(
                        month=ym, description=f"{desc} {res['p'][0]}/{res['p'][1]}", amount=abs(amt),
                        category='Carro - documentação' if 'seguro' in norm(desc) else 'Outros', where=where))
                    continue
                if kind == 'pix':
                    pix_items.append(dict(month=ym, description=desc, amount=abs(amt), where=where))
                    continue
                if res['credit']:
                    credits.append(dict(card=target, month=ym, amount=abs(amt), description=desc,
                                        anchor=re.sub(r'(?i)^estorno\s*', '', desc), where=where))
                    continue
                if amt < 0:
                    credits.append(dict(card=target, month=ym, amount=abs(amt), description=desc,
                                        anchor=None, where=where))
                    continue
                if res['p']:
                    raw_plan_items.append(dict(card=target, month=ym, description=desc, amount=amt,
                                               n=res['p'][0], total=res['p'][1], where=where))
                    continue
                card_items.append(dict(card=target, month=ym, description=desc, amount=amt, where=where))
                if res['estornado']:
                    credits.append(dict(card=target, month=ym, amount=amt, description=desc + ' (estorno)',
                                        anchor=desc, where=where))
        # fixos (histórico até o mês corrente)
        for col, (desc, cat, day) in sh['fixed'].items():
            v = ws[f'{col}{r}'].value
            if v and ym <= TODAY_MONTH:
                fixed_rows.append(dict(month=ym, description=desc, category=cat,
                                       amount=to_cents_float(v), dueDay=day))
        # receber (inclui projeção)
        v = ws[f"{sh['receber']}{r}"].value
        if v:
            receber.append(dict(month=ym, amount=abs(to_cents_float(v)), description='Receber (Larissa e outros)',
                                category='Larissa (entrada)'))
        # entradas (apenas até o mês corrente; depois vira recorrência de salário)
        c = ws[f"{sh['entradas']}{r}"]
        if c.value is None or ym > TODAY_MONTH:
            continue
        total = to_cents_float(c.value)
        items = []
        unknown = []
        where = f"{sh['name']}!{sh['entradas']}{r}"
        for ln in (body(c).split('\n') if c.comment else []):
            ln = ln.strip().replace('\xa0', ' ')
            if not ln:
                continue
            mm = re.match(r'^\+\s*(.*)$', ln)
            if not mm:
                flag('entrada: linha sem "+" (ignorada)', where, ln)
                continue
            txt = mm[1]
            nums = NUM.findall(re.sub(r'\b\d+/\d+\b', '', txt))
            if not nums:
                unknown.append(clean_desc(txt))
                continue
            amt = to_cents_str(pick_amount(nums, txt))
            d = clean_desc(txt) or 'Entrada'
            if norm(d) == 'parcela':
                d = '13º salário (1ª parcela)'
            nd = norm(d)
            cat = 'Saldo anterior' if 'saldo anterior' in nd else ('Salário' if 'salario' in nd else 'Extras')
            items.append(dict(month=ym, description=d, amount=amt, category=cat))
        rem = total - sum(i['amount'] for i in items)
        if len(unknown) == 1 and rem > SALARY_FUTURE_CENTS:
            inferred = rem - SALARY_FUTURE_CENTS
            items.append(dict(month=ym, description=unknown[0], amount=inferred, category='Extras'))
            rem -= inferred
            flag('valor inferido (célula − itens − salário padrão)', where, f'{unknown[0]} = {inferred / 100:.2f}')
        else:
            for u in unknown:
                flag('entrada sem valor (ignorada)', where, u)
        if rem > 0:
            items.append(dict(month=ym, description='Salário', amount=rem, category='Salário'))
        elif rem < 0:
            flag('itens de entrada somam mais que a célula', where, f'{rem / 100:.2f}')
        incomes.extend(items)

# ---------------------------------------------------------------- planos de parcelamento
def stem(d):
    d = norm(d)
    return re.sub(r'\bmagr[ao]s+\b', 'magras', d)


groups = collections.defaultdict(list)
for it in raw_plan_items:
    groups[(it['card'], stem(it['description']))].append(it)

plans = []
for (card, st), items in groups.items():
    items.sort(key=lambda x: x['month'])
    clusters = []
    for it in items:
        start = ym_index(it['month']) - it['n'] + 1
        placed = False
        for cl in clusters:
            if abs(start - cl['start']) <= 3 and it['month'] not in {i['month'] for i in cl['items']}:
                cl['items'].append(it)
                placed = True
                cnt = collections.Counter(ym_index(i['month']) - i['n'] + 1 for i in cl['items'])
                cl['start'] = max(cnt.items(), key=lambda kv: (kv[1], -kv[0]))[0]
                break
        if not placed:
            clusters.append(dict(items=[it], start=start))
    for cl in clusters:
        its = cl['items']
        tot_cnt = collections.Counter(i['total'] for i in its)
        total = max(tot_cnt.items(), key=lambda kv: (kv[1], -kv[0]))[0]
        start_i = cl['start']
        obs = {i['month']: i['amount'] for i in its}
        for i in its:
            derived = ym_index(i['month']) - start_i + 1
            if derived != i['n'] or i['total'] != total:
                flag('parcela corrigida pela continuidade', i['where'],
                     f"«{i['description']} {i['n']}/{i['total']}» → {derived}/{total}")
        first_obs = min(obs)
        if ym_index(max(obs)) > start_i + total - 1:
            flag('plano observado além da última parcela', 'plano', f'{its[0]["description"]} ({max(obs)})')
        desc = max(collections.Counter(i['description'] for i in its).items(), key=lambda kv: kv[1])[0]
        desc = desc[0].upper() + desc[1:]
        insts = []
        prev = obs[first_obs]
        for k in range(total):
            mi = ym_from_index(start_i + k)
            if mi in obs:
                prev = obs[mi]
                src = 'planilha'
            elif mi < first_obs:
                prev = obs[first_obs]
                src = 'estimada'
            else:
                src = 'estimada'
            insts.append(dict(number=k + 1, month=mi, amount=prev, source=src))
        plans.append(dict(card=card, description=desc, category=category_for(desc), count=total, declaredTotal=total,
                          startMonth=ym_from_index(start_i), installments=insts,
                          totalAmount=sum(x['amount'] for x in insts)))

# parcelas finais pagas por PIX (ex.: seguro 9/10 e 10/10 "saiu do cartão"): o plano de cartão termina antes
for pi in sorted(plain_installments, key=lambda x: -int(re.search(r'(\d+)/\d+$', x['description'])[1])):
    m = re.match(r'^(.*) (\d+)/(\d+)$', pi['description'])
    if not m:
        continue
    for pl in plans:
        if stem(pl['description']) == stem(m[1]) and pl['declaredTotal'] == int(m[3]):
            kept = [x for x in pl['installments'] if x['month'] != pi['month']]
            if len(kept) == len(pl['installments']) or any(x['number'] > int(m[2]) for x in kept):
                continue
            flag('parcela paga por PIX removida do plano de cartão', pi['where'],
                 f"{pl['description']} {m[2]}/{m[3]} em {pi['month']}")
            pl['installments'] = kept
            pl['count'] = len(kept)
            pl['totalAmount'] = sum(x['amount'] for x in kept)

# compra avulsa sem "n/total" que é, na verdade, a parcela 1 de um plano (ex.: "Bebel 162,50" e depois "Bebel 2/4")
for pl in plans:
    for x in pl['installments']:
        if x['source'] != 'estimada' or x['month'] > TODAY_MONTH:
            continue
        for ci in list(card_items):
            if ci['card'] == pl['card'] and ci['month'] == x['month'] and stem(ci['description']) == stem(pl['description']) \
                    and abs(ci['amount'] - x['amount']) <= max(100, x['amount'] // 50):
                card_items.remove(ci)
                x['source'] = 'planilha (sem n/total)'
                x['amount'] = ci['amount']
                flag('compra avulsa tratada como parcela 1 do plano', ci['where'], f"{ci['description']} → {pl['description']} 1/{pl['count']}")
                break

for pl in plans:
    pl['description'] = re.sub(r'^Ml\b', 'ML', pl['description'])
    pl['totalAmount'] = sum(x['amount'] for x in pl['installments'])

# ---------------------------------------------------------------- avulsos só até o mês corrente
card_items = [i for i in card_items if i['month'] <= TODAY_MONTH]
credits = [c for c in credits if c['month'] <= TODAY_MONTH]
pix_items = [i for i in pix_items if i['month'] <= TODAY_MONTH]
for i in card_items + pix_items:
    i['category'] = category_for(i['description'])

# ---------------------------------------------------------------- recorrências
first_future = add_months(TODAY_MONTH, 1)
ws1 = wb['AGO26-ATUAL']
recurrences = []
for desc, v in FUTURE_FIXED.items():
    recurrences.append(dict(description=desc, amount=v['amount'], category=v['category'], method='PIX',
                            startMonth=first_future, endMonth=None, dueDay=v['dueDay'], type='EXPENSE'))
recurrences.append(dict(description='Salário', amount=SALARY_FUTURE_CENTS, category='Salário', method='PIX',
                        startMonth=first_future, endMonth=None, dueDay=1, type='INCOME'))
subs = {}
for r in SHEETS[1]['rows']:
    ym = month_of(ws1, r)
    if ym < first_future:
        continue
    c = ws1[f'J{r}']
    if not c.comment:
        continue
    for ln in body(c).split('\n'):
        res = parse_line(ln)
        if not res or res['p'] or not res['nums']:
            continue
        d = clean_desc(res['rest'])
        if norm(d) in SUBSCRIPTION_NAMES or norm(d).startswith('wellhub'):
            e = subs.setdefault(norm(d), dict(description=re.sub(r' x\d$', '', d).capitalize() if d.islower() else d, amounts={}))
            e['amounts'][ym] = to_cents_str(pick_amount(res['nums'], res['raw']))
last_month = max(months_seen)
for key, e in subs.items():
    months = sorted(e['amounts'])
    last = months[-1]
    recurrences.append(dict(description=e['description'], amount=e['amounts'][last], category='Assinaturas',
                            method='Bradesco Amazon', startMonth=months[0],
                            endMonth=None if last == last_month else last, dueDay=1, type='EXPENSE'))
    if ym_index(last) - ym_index(months[0]) + 1 != len(months):
        flag('assinatura com lacunas entre meses futuros', 'assinaturas', e['description'])

# ---------------------------------------------------------------- dataset
dataset = dict(
    todayMonth=TODAY_MONTH,
    cards=[dict(name=k, **v) for k, v in CARDS.items()],
    categories=[dict(name=n, type=t) for n, t in CATEGORIES],
    cardPurchases=card_items, credits=credits, plans=plans,
    plainInstallments=plain_installments, pixItems=pix_items,
    fixed=fixed_rows, receber=receber, incomes=incomes, recurrences=recurrences,
    sheetCells={f'{k[0]}|{k[1]}': v for k, v in sheet_cells.items()},
)
with open(os.path.join(OUT, 'dataset.json'), 'w', encoding='utf-8') as f:
    json.dump(dataset, f, ensure_ascii=False, indent=1)


# ---------------------------------------------------------------- relatório
def money(c):
    return f'{c / 100:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


L = ['# Revisão da importação (somente leitura, nada foi gravado)\n']
L.append(f'- Compras de cartão avulsas (jan–out/26): **{len(card_items)}**')
L.append(f'- Planos de parcelamento: **{len(plans)}** — parcelas: **{sum(len(p["installments"]) for p in plans)}**')
L.append(f'  - anteriores a {FIRST_MONTH} (criadas como pagas): '
         f'{sum(1 for p in plans for x in p["installments"] if x["month"] < FIRST_MONTH)}')
L.append(f'  - futuras (a partir de {first_future}): '
         f'{sum(1 for p in plans for x in p["installments"] if x["month"] >= first_future)}')
L.append(f'- Débito/PIX avulsos: **{len(pix_items)}**; parcelas comuns (celular/seguro): **{len(plain_installments)}**')
L.append(f'- Créditos/estornos: **{len(credits)}**')
L.append(f'- Fixos jan–out: **{len(fixed_rows)}**; Receber: **{len(receber)}** meses; '
         f'Entradas: **{len(incomes)}** lançamentos')
L.append(f'- Recorrências (a partir de {first_future}): **{len(recurrences)}**\n')

L.append('## Fatura: valor da célula x total importado (jan–out/26)\n')
L.append('| Cartão | Mês | Célula | Importado | Diferença |')
L.append('|---|---|---:|---:|---:|')
tot = collections.defaultdict(int)
for i in card_items:
    tot[(i['card'], i['month'])] += i['amount']
for p in plans:
    for x in p['installments']:
        if x['month'] <= TODAY_MONTH:
            tot[(p['card'], x['month'])] += x['amount']
for c in credits:
    tot[(c['card'], c['month'])] -= c['amount']
for (card, ym), cell in sorted(sheet_cells.items(), key=lambda kv: (kv[0][1], kv[0][0])):
    if ym > TODAY_MONTH:
        continue
    got = tot.get((card, ym), 0)
    d = got - cell
    L.append(f'| {card} | {ym} | {money(cell)} | {money(got)} | {money(d)}{" ⚠" if abs(d) > 5000 else ""} |')
L.append('')

L.append('## Planos de parcelamento\n')
L.append('| Descrição | Cartão | Categoria | Início | Parcelas | Total |')
L.append('|---|---|---|---|---:|---:|')
for p in sorted(plans, key=lambda p: (p['startMonth'], p['description'])):
    L.append(f'| {p["description"]} | {p["card"]} | {p["category"]} | {p["startMonth"]} | {p["count"]} | '
             f'{money(p["totalAmount"])} |')
L.append('')

L.append('## Itens em "Outros" (revise a categoria)\n')
cnt = collections.Counter(norm(i['description']) for i in card_items + pix_items + plans if i['category'] == 'Outros')
for d, n in cnt.most_common():
    L.append(f'- {d} ({n}x)')
L.append('')

L.append('## Sinalizações\n')
for kind, where, text in flags:
    L.append(f'- **{kind}** — {where}: `{text}`')
L.append('')

L.append('## Recorrências\n')
for r in recurrences:
    L.append(f'- {r["description"]}: {money(r["amount"])} ({r["type"]}, {r["category"]}, {r["method"]}) '
             f'início {r["startMonth"]}' + (f' fim {r["endMonth"]}' if r['endMonth'] else ''))
L.append('')

L.append('## Entradas jan–out\n')
for i in sorted(incomes, key=lambda x: x['month']):
    L.append(f'- {i["month"]} {i["category"]}: {i["description"]} {money(i["amount"])}')
L.append('')

with io.open(os.path.join(OUT, 'review.md'), 'w', encoding='utf-8') as f:
    f.write('\n'.join(L))
print('ok:', OUT)
print(f'{len(card_items)} avulsos cartão, {len(plans)} planos, {len(pix_items)} pix, {len(credits)} créditos, '
      f'{len(flags)} sinalizações')
